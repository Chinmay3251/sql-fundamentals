import io
from fastapi.testclient import TestClient
from pptx import Presentation
import app.main as main
client=TestClient(main.app)
CSV=b"region,sales,orders\nNorth,100,4\nSouth,200,5\nNorth,150,3\nEast,50,2\n"
def test_eda_returns_five_insights():
 r=client.post("/api/analyze",files={"file":("sales.csv",CSV,"text/csv")}); assert r.status_code==200; assert r.json()["rows"]==4 and len(r.json()["insights"])==5
def test_summary_200_word_limit(monkeypatch):
 monkeypatch.delenv("GROQ_API_KEY",raising=False); r=client.post("/api/summary",json={"insights":[f"Insight {i}: sales pattern." for i in range(5)]}); assert r.status_code==200 and r.json()["word_count"]<=200
def test_ppt_seven_slides():
 r=client.post("/api/export-ppt",json={"title":"Test brief","insights":[f"Insight {i}" for i in range(1,6)],"recommendations":["Validate data"],"summary":"Test summary"}); assert r.status_code==200; prs=Presentation(io.BytesIO(r.content)); assert len(prs.slides)==7 and prs.slides[-1].shapes.title.text=="Recommendations"
def test_full_flow_csv_to_ppt(monkeypatch):
 monkeypatch.setattr(main,"generate_summary",lambda insights:{"summary":"Mock executive summary.","provider":"mock","word_count":3}); r=client.post("/api/full-flow",files={"file":("sales.csv",CSV,"text/csv")}); assert r.status_code==200; assert len(Presentation(io.BytesIO(r.content)).slides)==7
def test_reject_non_csv(): assert client.post("/api/full-flow",files={"file":("notes.txt",b"hello","text/plain")}).status_code==400
