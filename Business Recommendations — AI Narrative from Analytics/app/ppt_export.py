from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt

def build_presentation(title, insights, recommendations, summary=""):
    if len(insights)!=5: raise ValueError("PowerPoint export requires exactly five insights.")
    if not recommendations: recommendations=["Validate data quality and metric definitions.","Investigate material trends and outliers.","Assign owners and monitor KPIs regularly."]
    prs=Presentation(); prs.core_properties.title=title; prs.core_properties.subject="EDA executive summary"; prs.core_properties.author="CIA Analytics"
    slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text=title; slide.placeholders[1].text="Exploratory Data Analysis | Executive Brief"
    if summary:
        box=slide.shapes.add_textbox(Inches(.7),Inches(5.2),Inches(8),Inches(1.2)); box.text_frame.text=summary[:700]; box.text_frame.paragraphs[0].font.size=Pt(12)
    for i,insight in enumerate(insights,1):
        slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text=f"Insight {i}"
        box=slide.shapes.add_textbox(Inches(.8),Inches(1.7),Inches(8.2),Inches(3.7)); box.text_frame.word_wrap=True; p=box.text_frame.paragraphs[0]; p.text=insight; p.font.size=Pt(24)
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text="Recommendations"
    box=slide.shapes.add_textbox(Inches(.8),Inches(1.5),Inches(8.2),Inches(4.8)); tf=box.text_frame; tf.clear()
    for i,rec in enumerate(recommendations[:8]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text="• "+rec; p.font.size=Pt(20); p.space_after=Pt(14)
    out=BytesIO(); prs.save(out); return out.getvalue()
