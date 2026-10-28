from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt

def add_bullets(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    frame = slide.placeholders[1].text_frame; frame.clear()
    for i, item in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = str(item); p.font.size = Pt(20); p.space_after = Pt(14)
    return slide

def create_presentation(filename, rows, columns, insights, summary, recommendations):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = 'Business Recommendations'
    slide.placeholders[1].text = f'AI Narrative from Analytics\nSource: {filename}\nRecords: {rows:,} | Columns: {len(columns):,}'
    for i, insight in enumerate(insights[:5], 1): add_bullets(prs, f'Insight {i}', [insight])
    add_bullets(prs, 'Recommendations', [f'Executive summary: {summary}', *recommendations[:3]])
    out = BytesIO(); prs.save(out); return out.getvalue()
